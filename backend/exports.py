import io,json,re
import pandas as pd
from openpyxl.styles import Font,PatternFill,Alignment
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4,landscape
from reportlab.platypus import SimpleDocTemplate,Paragraph,Spacer,Table,TableStyle
from reportlab.lib.styles import getSampleStyleSheet
from xml.sax.saxutils import escape
from storage import folder,read

def safe(v):
    if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@','\t','\r')):return "'"+v
    return v

def tabular(df,fmt):
    df=df.copy();df.columns=[safe(str(c)) for c in df.columns];df=df.map(safe);buf=io.BytesIO()
    if fmt=='csv':return df.to_csv(index=False,na_rep='').encode('utf-8-sig'),'text/csv'
    with pd.ExcelWriter(buf,engine='openpyxl') as writer:
        df.to_excel(writer,index=False,sheet_name='LeadIQ results',na_rep='');ws=writer.sheets['LeadIQ results'];ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='6355D8');cell.alignment=Alignment(wrap_text=True)
        for col in ws.columns:
            width=min(42,max(12,max(len(str(c.value or '')) for c in list(col)[:50])+2));ws.column_dimensions[col[0].column_letter].width=width
        ws.row_dimensions[1].height=30
    return buf.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

def methodology(pid):
    m=read(pid,'metrics');s=read(pid,'score_summary');best=m['best'];f=read(pid,'features')
    text=(f'LeadIQ trained five binary-classification candidates using random seed 42. The original stratified 80/20 split was preserved through structural cleaning, leaving {m["split"]["train"]} training and {m["split"]["test"]} held-out test rows. All learned transformations were fitted within training folds. {m["best_model"]} was selected by {m["folds"]}-fold training cross-validated average precision, then calibrated using {m["calibration_method"]}. The decision threshold ({m["threshold"]:.4f}) maximized F1 on out-of-fold calibrated training predictions. Held-out ROC-AUC was {best["roc_auc"]:.4f}, average precision (PR-AUC) {best["pr_auc"]:.4f}, and F1 {best["f1"]:.4f}. Confidence intervals used 1,000 bootstrap resamples. Lead Scores rank calibrated probabilities against the training probability distribution; they are not conversion probabilities. {s["total"]} {s["source"]} records were scored. SHAP explains the underlying uncalibrated estimator, not causation. Dataset SHA-256: {m["dataset_hash"]}.')
    return {'methodology':text,'best_model':m['best_model'],'scored_rows':s['total'],'features':f['count'],'metrics':best,'seed':42,'dataset_hash':m['dataset_hash'],'versions':m['versions'],'parameters':{r['name']:r['parameters'] for r in m['models']},'calibration':m['calibration_method'],'threshold':m['threshold'],'split':m['split'],'warnings':m['warnings']}

def pdf_export(pid,top_n):
    summary=methodology(pid);m=read(pid,'metrics');df=pd.read_pickle(folder(pid)/'scores.pkl').head(top_n);buf=io.BytesIO();doc=SimpleDocTemplate(buf,pagesize=landscape(A4),rightMargin=30,leftMargin=30,topMargin=30,bottomMargin=30);styles=getSampleStyleSheet();story=[]
    def p(t,style='BodyText'):return Paragraph(escape(str(t)),styles[style])
    story.extend([p('LeadIQ | Run summary','Title'),Spacer(1,12),p(summary['methodology']),Spacer(1,16),p('Model comparison','Heading2')])
    rows=[['Model','ROC-AUC','PR-AUC','F1','Precision@10%']]+[[r['name']]+[f'{r[k]:.4f}' for k in ['roc_auc','pr_auc','f1','precision_top10']] for r in m['models']]
    table=Table(rows,hAlign='LEFT');table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#6355d8')),('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.3,colors.lightgrey),('PADDING',(0,0),(-1,-1),7)]));story.extend([table,Spacer(1,16),p(f'Top {len(df)} ranked records','Heading2')])
    cfg=read(pid,'config');name=next((c for c,r in cfg['roles'].items() if r=='Business Name' and c in df),'lead_id')
    rows=[['Order','Record','Score','Probability','Tier']]+[[str(r.contact_order),p(str(r[name])[:90]),f'{r.lead_score:.2f}',f'{r.probability:.4f}',r.tier] for _,r in df.iterrows()]
    table=Table(rows,colWidths=[45,350,70,90,95],repeatRows=1,hAlign='LEFT');table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#eeeafa')),('GRID',(0,0),(-1,-1),.2,colors.lightgrey),('PADDING',(0,0),(-1,-1),6)]));story.append(table)
    story.extend([Spacer(1,16),p('Cautions','Heading2'),p('Scores are prioritization aids, not guarantees. Handle personal information responsibly. Projects expire after 24 hours.')])
    for w in m['warnings']:story.append(p(w))
    doc.build(story);return buf.getvalue(),'application/pdf'